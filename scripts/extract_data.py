import arcpy
import json
import openpyxl

sr_wgs84 = arcpy.SpatialReference(4326)

# 1. Read ESTACIONES
est_fields = [f.name for f in arcpy.ListFields('SHP/ESTACIONES.shp')]
print("ESTACIONES fields:", est_fields)

estaciones_data = []
with arcpy.da.SearchCursor('SHP/ESTACIONES.shp', ['SHAPE@XY'] + est_fields, spatial_reference=sr_wgs84) as cursor:
    for row in cursor:
        xy = row[0]
        attr = dict(zip(est_fields, row[1:]))
        attr['lon'] = xy[0]
        attr['lat'] = xy[1]
        estaciones_data.append(attr)

print(f"Read {len(estaciones_data)} estaciones:")
for e in estaciones_data:
    print(e.get('ESTACION'), e.get('NOMBRE_EST'), e.get('COMPAÑIA') or e.get('COMPANIA'), f"({e['lat']:.5f}, {e['lon']:.5f})")

# 2. Read Localidades
loc_fields = [f.name for f in arcpy.ListFields('SHP/localidades.shp')]
print("\nLOCALIDADES fields:", loc_fields)

localidades_data = []
with arcpy.da.SearchCursor('SHP/localidades.shp', ['SHAPE@JSON'] + loc_fields, spatial_reference=sr_wgs84) as cursor:
    for row in cursor:
        geom_json = json.loads(row[0])
        attr = dict(zip(loc_fields, row[1:]))
        localidades_data.append({
            'attr': attr,
            'geometry': geom_json
        })
print(f"Read {len(localidades_data)} localidades.")
for loc in localidades_data:
    print(loc['attr'].get('LocNombre'), loc['attr'].get('LocCodigo'), loc['attr'].get('SHAPE_Leng'))

# 3. Read Excel data
wb = openpyxl.load_workbook('2.1 BASE PREDIAL ESTACIONES - ACTUALIZADA SEPTIEMBRE DIRECCIONES CONFIRMADAS.xlsx', data_only=True)
sheet = wb['BASE']
headers = [str(sheet.cell(1, c).value).strip() for c in range(1, sheet.max_column + 1)]
excel_rows = []
for r in range(2, sheet.max_row + 1):
    vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
    if any(vals):
        row_dict = {headers[i]: vals[i] for i in range(len(headers))}
        excel_rows.append(row_dict)

with open('scratch_estaciones.json', 'w', encoding='utf-8') as f:
    json.dump({'estaciones_shp': estaciones_data, 'excel': excel_rows}, f, indent=2, ensure_ascii=False, default=str)

print("Saved scratch_estaciones.json")
