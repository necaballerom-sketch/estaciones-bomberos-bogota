import json
import shutil
import os
import openpyxl
import arcpy

sr_wgs84 = arcpy.SpatialReference(4326)

# 1. Read shapefile stations
shp_stations = {}
fields = [f.name for f in arcpy.ListFields('SHP/ESTACIONES.shp')]
with arcpy.da.SearchCursor('SHP/ESTACIONES.shp', ['SHAPE@XY'] + fields, spatial_reference=sr_wgs84) as cursor:
    for row in cursor:
        xy = row[0]
        attr = dict(zip(fields, row[1:]))
        code = (attr.get('ESTACION') or '').strip()
        shp_stations[code] = {
            'code': code,
            'nombre_shp': (attr.get('NOMBRE_EST') or '').strip(),
            'direccion_shp': (attr.get('DIRECCION') or '').strip(),
            'telefonos_shp': (attr.get('TELEFONOS') or '').strip(),
            'compania_shp': (attr.get('COMPAÑIA') or attr.get('COMPA\x91IA') or attr.get('COMPAÑIA') or '').strip(),
            'lat': round(xy[1], 6),
            'lon': round(xy[0], 6)
        }

# 2. Read Excel
wb = openpyxl.load_workbook('2.1 BASE PREDIAL ESTACIONES - ACTUALIZADA SEPTIEMBRE DIRECCIONES CONFIRMADAS.xlsx', data_only=True)
sheet = wb['BASE']
headers = [str(sheet.cell(1, c).value).strip() for c in range(1, sheet.max_column + 1)]
excel_data = {}
for r in range(2, sheet.max_row + 1):
    vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
    if any(vals):
        code_raw = str(vals[0]).strip() if vals[0] else ''
        name_raw = str(vals[1]).strip() if vals[1] else ''
        row_dict = {headers[i]: vals[i] for i in range(len(headers))}
        excel_data[code_raw] = row_dict

print("SHP stations keys:", list(shp_stations.keys()))
print("Excel stations keys:", list(excel_data.keys()))

# 3. Handle photo files & aliases
fotos_dir = 'FOTOS'
# Create clean alias for B20 LA CABAÑA.jpg
for fn in os.listdir(fotos_dir):
    if 'CABA' in fn:
        src = os.path.join(fotos_dir, fn)
        dst1 = os.path.join(fotos_dir, 'B20_LA_CABANA.jpg')
        dst2 = os.path.join(fotos_dir, 'B20 LA CABANA.jpg')
        if not os.path.exists(dst1):
            shutil.copy2(src, dst1)
        if not os.path.exists(dst2):
            shutil.copy2(src, dst2)
        print("Created alias for Cabaña photo")

print("Done Step 1 & 2")
